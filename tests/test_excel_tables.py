"""Tests for structured Excel table extraction."""

import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from store_brief.parse.excel_tables import (
    extract_tables_from_sheet,
    extract_tables_from_workbook,
    tables_to_text,
)


def test_simple_table():
    df = pd.DataFrame(
        [
            ["상품", "재고", "가격"],
            ["모델A", 10, 1000],
            ["모델B", 5, 2000],
        ]
    )
    tables = extract_tables_from_sheet("재고", df)
    assert len(tables) == 1
    table = tables[0]
    assert table.columns == ["상품", "재고", "가격"]
    assert len(table.rows) == 2
    assert table.rows[0][0] == "모델A"


def test_title_and_skip_blank_rows():
    df = pd.DataFrame(
        [
            ["▣ 월간 재고 현황", None, None],
            [None, None, None],
            ["구분", "수량", "금액"],
            ["합계", 100, 500],
            ["A", 40, 200],
        ]
    )
    tables = extract_tables_from_sheet("현황", df)
    assert len(tables) == 1
    table = tables[0]
    assert "재고" in (table.title or "")
    assert table.columns[0] == "구분"
    assert table.rows[-1][0] == "A"


def test_multi_row_header_merge():
    df = pd.DataFrame(
        [
            ["▣ 지사별 소진 현황", None, None, None],
            [None, None, None, None],
            ["구분", "기초", None, "소진율"],
            [None, "3월말", "5월말", None],
            ["합계", 100, 80, 0.2],
            ["강북", 10, 8, 0.2],
        ]
    )
    tables = extract_tables_from_sheet("현황", df)
    assert len(tables) == 1
    table = tables[0]
    assert "구분" in table.columns[0]
    assert "기초" in table.columns[1]
    assert "3월말" in table.columns[1]
    assert table.rows[0][0] == "합계"


def test_multiple_regions_in_one_sheet():
    df = pd.DataFrame(
        [
            ["▣ 클리어런스", None, None],
            ["모델", "재고", "가격"],
            ["A", 10, 1000],
            [None, None, None],
            [None, None, None],
            ["▣ 온라인", None, None],
            ["모델", "판매", "비고"],
            ["B", 5, "쿠팡"],
        ]
    )
    tables = extract_tables_from_sheet("mixed", df)
    assert len(tables) == 2
    assert tables[0].rows[0][0] == "A"
    assert tables[1].rows[0][0] == "B"
    assert tables[1].region == 1


def test_workbook_and_markdown():
    sheets = {
        "s1": pd.DataFrame([["a", "b"], [1, 2]]),
    }
    tables = extract_tables_from_workbook(sheets)
    assert len(tables) == 1
    text = tables_to_text(tables)
    assert "### 시트: s1" in text
    assert "| a | b |" in text


def test_simple_list_many_columns():
    """Wide header row must not be mistaken for a title."""
    df = pd.DataFrame(
        [
            ["", "진열상품코드", "상품명", "SET유형", "품목코드", "품목명", "담당부서명", "운영유형"],
            ["1", "(J)ABC", "테스트 상품", "일반상품", "EAFCA", "키트", "에어컨팀", "전점"],
        ]
    )
    tables = extract_tables_from_sheet("전시소진", df)
    assert len(tables) == 1
    assert tables[0].columns[0] == "No."
    assert tables[0].columns[1] == "진열상품코드"
    assert tables[0].rows[0][1] == "(J)ABC"


def test_excel_parse_reads_workbook_once(monkeypatch, tmp_path):
    import openpyxl
    from types import SimpleNamespace

    from store_brief.parse import excel as excel_mod

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["No.", "상품"])
    ws.append(["1", "A"])
    path = tmp_path / "sample.xlsx"
    wb.save(path)

    calls: list[int] = []
    real_read = excel_mod.read_workbook

    def counting_read(p):
        calls.append(1)
        return real_read(p)

    monkeypatch.setattr(excel_mod, "read_workbook", counting_read)
    monkeypatch.setattr(excel_mod, "rasterize_or_empty", lambda *a, **k: [])
    att = SimpleNamespace(id="att1", path=str(path), filename="sample.xlsx")
    excel_mod.parse(att, str(tmp_path))

    assert len(calls) == 1


def test_excel_parse_skips_com_images_without_vlm(monkeypatch, tmp_path):
    import openpyxl
    from types import SimpleNamespace

    from store_brief.parse import excel as excel_mod

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "sheet1"
    ws.append(["No.", "상품"])
    ws.append(["1", "A"])
    path = tmp_path / "sample.xlsx"
    wb.save(path)

    called: list[bool] = []

    def _boom(*_args, **_kwargs):
        called.append(True)
        return []

    monkeypatch.setattr(excel_mod, "rasterize_or_empty", _boom)
    att = SimpleNamespace(id="att1", path=str(path), filename="sample.xlsx")
    result = excel_mod.parse(att, str(tmp_path))

    assert not called
    assert result.image_paths == []
    assert result.raw_sheets
    assert result.tables


if __name__ == "__main__":
    test_simple_table()
    test_title_and_skip_blank_rows()
    test_multi_row_header_merge()
    test_multiple_regions_in_one_sheet()
    test_workbook_and_markdown()
    print("ok")
